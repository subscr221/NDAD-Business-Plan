#!/usr/bin/env python3
"""
build_v4_docx_pdf.py - Regenerate the finalized investor-plan deliverables
(Nitrodynamics_Business_Plan_V4_final.docx + .pdf) from the consolidated
markdown master, with the 10 diagram PNGs re-embedded by section.

Pipeline note: the ORIGINAL .docx/.pdf were produced by LibreOffice-on-Linux,
which is not present on this host. This script is the supported replacement -
it derives both deliverables directly from investor_plan/Nitrodynamics_Investor_Plan.md
(the true source of truth, itself rebuilt by build_investor_plan.py from sections 00-10).

Run order after editing any section file:
    python build_investor_plan.py      # rebuild the markdown master
    python build_v4_docx_pdf.py        # rebuild docx + pdf from the master

Engines: python-docx (Word) + fpdf2 (PDF, Arial TTF for full Unicode).
Requires: pip install python-docx fpdf2   (Windows Arial fonts for the PDF).
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from fpdf import FPDF
from fpdf.enums import XPos, YPos

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "investor_plan"
DIAG = SRC / "diagrams"
MD = SRC / "Nitrodynamics_Investor_Plan.md"

OUT_DOCX = SRC / "Nitrodynamics_Business_Plan_V4_final.docx"
OUT_PDF = SRC / "Nitrodynamics_Business_Plan_V4_final.pdf"

FONTS = Path("C:/Windows/Fonts")  # Arial TTFs (Unicode); change if not on Windows.

# ---- diagram registry (number -> file, caption) per the markdown Diagram Index ----
DIAGRAMS = {
    1: ("01_platform_flywheel.png", "Platform Flywheel - M1-M5 modules compounding into V1-V6 verticals"),
    2: ("02_roadmap_gantt.png", "10-Year Roadmap Gantt - module & product maturity by year, funding milestones"),
    3: ("03_revenue_build.png", "Revenue Build by Product - stacked Y1-Y10 contribution by vertical"),
    4: ("04_competitive_scorecard.png", "Competitive Scorecard Heatmap - 12 capabilities x 7 archetypes"),
    5: ("05_three_phase_arc.png", "Three-Phase Strategic Arc - insurgent > productised scale-up > multi-domain prime"),
    6: ("06_working_capital_cycle.png", "Working Capital Cycle - 225-day CCC, BG collateral schedule"),
    7: ("07_funding_stack.png", "Funding Stack vs Cash Curve - equity/debt slabs vs deployable-cash polyline"),
    8: ("08_operating_org.png", "Operating Model Org Chart - Board, CEO, five directs, platform & product squads"),
    9: ("09_pnl_waterfall.png", "P&L Waterfall (Year 10) - Revenue Rs.5,055 Cr to Net Income"),
    10: ("10_sensitivity_tornado.png", "Sensitivity Tornado - eight single-variable swings on IRR (base 31.9%)"),
}
# H1-title substring -> list of diagram numbers placed after that section's heading.
# Order matters: more specific keys first ("Strategy, Operating Model" before "Operating Model").
SECTION_DIAGRAMS = {
    "Strategy, Operating Model": [1, 2, 4, 5],
    "Financial Model": [3, 6, 9],
    "Funding Plan": [7],
    "Returns, Valuation": [10],
    "Operating Model": [8],
}

# ----------------------------- markdown tokenizer -----------------------------
def is_block_start(line, nxt):
    s = line.strip()
    if s == "":
        return True
    if s == "<!-- page-break -->" or s == "---":
        return True
    if re.match(r"^#{1,6}\s+", line):
        return True
    if s.startswith("```") or s.startswith(">"):
        return True
    if re.match(r"^\s*[-*]\s+", line) or re.match(r"^\s*\d+\.\s+", line):
        return True
    if s.startswith("|") and re.match(r"^\s*\|[\s:|-]+\|\s*$", nxt):
        return True
    return False


def tokenize(md):
    lines = md.split("\n")
    n = len(lines)
    toks = []
    i = 0
    while i < n:
        line = lines[i]
        s = line.strip()
        nxt = lines[i + 1] if i + 1 < n else ""
        if s == "<!-- page-break -->":
            toks.append(("pagebreak", None)); i += 1; continue
        if s == "---":
            i += 1; continue  # section HR; pairs with the page-break marker
        if s.startswith("<!--") and s.endswith("-->"):
            i += 1; continue
        if s.startswith("```"):
            buf = []; i += 1
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            toks.append(("code", "\n".join(buf))); continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            toks.append(("h%d" % len(m.group(1)), m.group(2).strip())); i += 1; continue
        if s.startswith("|") and re.match(r"^\s*\|[\s:|-]+\|\s*$", nxt):
            buf = []
            while i < n and lines[i].strip().startswith("|"):
                buf.append(lines[i]); i += 1
            toks.append(("table", buf)); continue
        if re.match(r"^\s*[-*]\s+", line):
            items = []
            while i < n and re.match(r"^\s*[-*]\s+", lines[i]):
                indent = len(lines[i]) - len(lines[i].lstrip())
                items.append((indent, re.sub(r"^\s*[-*]\s+", "", lines[i]).rstrip())); i += 1
            toks.append(("ul", items)); continue
        if re.match(r"^\s*\d+\.\s+", line):
            items = []
            while i < n and re.match(r"^\s*\d+\.\s+", lines[i]):
                indent = len(lines[i]) - len(lines[i].lstrip())
                items.append((indent, re.sub(r"^\s*\d+\.\s+", "", lines[i]).rstrip())); i += 1
            toks.append(("ol", items)); continue
        if s.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i])); i += 1
            toks.append(("quote", " ".join(buf))); continue
        if s == "":
            i += 1; continue
        para = [line]; i += 1
        while i < n and lines[i].strip() != "" and not is_block_start(lines[i], lines[i + 1] if i + 1 < n else ""):
            para.append(lines[i]); i += 1
        toks.append(("p", " ".join(x.strip() for x in para)))
    return toks


def inline_segments(text):
    """Split into (text, styleset) where style in {b,i,code}."""
    pat = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*?\*)")
    out = []; pos = 0
    for m in pat.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], set()))
        t = m.group(0)
        if t.startswith("**"):
            out.append((t[2:-2], {"b"}))
        elif t.startswith("`"):
            out.append((t[1:-1], {"code"}))
        else:
            out.append((t[1:-1], {"i"}))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], set()))
    return out or [(text, set())]


def parse_table(buf):
    rows = []
    for ln in buf:
        s = ln.strip()
        if re.match(r"^\|[\s:|-]+\|$", s):
            continue
        rows.append([c.strip() for c in s.strip("|").split("|")])
    return rows


def png_size(path):
    with open(path, "rb") as f:
        data = f.read(33)
    return int.from_bytes(data[16:20], "big"), int.from_bytes(data[20:24], "big")


def diagrams_for_h1(title):
    for key, nums in SECTION_DIAGRAMS.items():
        if key in title:
            return nums
    return []


def clean_heading(text):
    return text.replace("**", "").replace("`", "").strip()


# ============================= DOCX RENDERER =============================
def shade_cell(cell, fill="D9E2F3"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_inline_docx(p, text, size=None, bold_all=False):
    for seg, st in inline_segments(text):
        if seg == "":
            continue
        r = p.add_run(seg)
        if "b" in st or bold_all:
            r.bold = True
        if "i" in st:
            r.italic = True
        if "code" in st:
            r.font.name = "Consolas"; r.font.size = Pt(9)
        elif size:
            r.font.size = size


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Nitrodynamics  -  Confidential  -  Page ").font.size = Pt(8)
    run = p.add_run()
    for kind, txt in (("begin", None), ("instr", "PAGE"), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        run._r.append(el)
    run.font.size = Pt(8)


def add_table_docx(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.autofit = True
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(ncol):
            cell = cells[ci]
            cell.text = ""
            add_inline_docx(cell.paragraphs[0], row[ci] if ci < len(row) else "",
                            size=Pt(9), bold_all=(ri == 0))
            if ri == 0:
                shade_cell(cell)
    doc.add_paragraph()


def insert_diagrams_docx(doc, nums):
    for num in nums:
        fname, caption = DIAGRAMS[num]
        path = DIAG / fname
        if not path.exists():
            print("  [warn] diagram missing:", path.name)
            continue
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(6.0))
        cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cap.add_run("Figure %d. %s" % (num, caption))
        cr.italic = True; cr.font.size = Pt(8.5); cr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def build_docx(tokens):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(10.5)
    add_page_number_footer(doc.sections[0])

    seen_first_h1 = False
    for typ, payload in tokens:
        if typ == "pagebreak":
            doc.add_page_break()
        elif typ.startswith("h"):
            level = int(typ[1])
            if level == 1 and not seen_first_h1:
                seen_first_h1 = True
                tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                tr = tp.add_run(clean_heading(payload)); tr.bold = True; tr.font.size = Pt(30)
                tr.font.color.rgb = RGBColor(0x1F, 0x33, 0x64)
                st = doc.add_paragraph(); st.alignment = WD_ALIGN_PARAGRAPH.CENTER
                sr = st.add_run("Defense Technology Platform  |  Business Plan  |  V4 (Final)")
                sr.font.size = Pt(13); sr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                continue
            doc.add_heading(clean_heading(payload), level=min(level, 4))
            if level == 1:
                insert_diagrams_docx(doc, diagrams_for_h1(payload))
        elif typ == "p":
            add_inline_docx(doc.add_paragraph(), payload)
        elif typ == "ul":
            for _, item in payload:
                add_inline_docx(doc.add_paragraph(style="List Bullet"), item)
        elif typ == "ol":
            for _, item in payload:
                add_inline_docx(doc.add_paragraph(style="List Number"), item)
        elif typ == "quote":
            add_inline_docx(doc.add_paragraph(style="Intense Quote"), payload)
        elif typ == "code":
            r = doc.add_paragraph().add_run(payload)
            r.font.name = "Consolas"; r.font.size = Pt(9)
        elif typ == "table":
            add_table_docx(doc, parse_table(payload))
    doc.save(str(OUT_DOCX))


# ============================= PDF RENDERER =============================
class PDF(FPDF):
    def footer(self):
        self.set_y(-12)
        self.set_font("Arial", "I", 8)
        self.set_text_color(120, 120, 120)
        self.cell(0, 8, "Nitrodynamics  -  Confidential  -  Page %d" % self.page_no(), align="C")

    def block(self, h, txt, **kw):
        """multi_cell that always starts at the left margin and advances down."""
        self.set_x(self.l_margin)
        kw.setdefault("new_x", XPos.LMARGIN)
        kw.setdefault("new_y", YPos.NEXT)
        return self.multi_cell(0, h, txt, **kw)


def md_for_pdf(text):
    """Keep **bold**; drop inline-code backticks (fpdf markdown has no code style)."""
    return text.replace("`", "")


def insert_diagrams_pdf(pdf, nums, epw):
    for num in nums:
        fname, caption = DIAGRAMS[num]
        path = DIAG / fname
        if not path.exists():
            continue
        w, h = png_size(path)
        disp_w = min(epw, 175)
        disp_h = disp_w * h / w
        if pdf.get_y() + disp_h + 12 > pdf.h - 15:
            pdf.add_page()
        pdf.image(str(path), x=(pdf.w - disp_w) / 2, w=disp_w)
        pdf.set_y(pdf.get_y() + disp_h + 1)
        pdf.set_font("Arial", "I", 8.5); pdf.set_text_color(85, 85, 85)
        pdf.block(4.5, "Figure %d. %s" % (num, caption), align="C")
        pdf.set_text_color(0, 0, 0); pdf.ln(3)


def render_table_pdf(pdf, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    pdf.set_font("Arial", "", 7.5)
    pdf.ln(1)
    with pdf.table(first_row_as_headings=True, markdown=True,
                   borders_layout="MINIMAL", line_height=4.2, text_align="LEFT") as table:
        for r in rows:
            row = table.row()
            for c in r:
                row.cell(md_for_pdf(c))
    pdf.ln(2)
    pdf.set_font("Arial", "", 10)


def build_pdf(tokens):
    pdf = PDF(orientation="P", unit="mm", format="A4")
    for style, fn in (("", "arial.ttf"), ("B", "arialbd.ttf"), ("I", "ariali.ttf"), ("BI", "arialbi.ttf")):
        f = FONTS / fn
        if not f.exists():
            sys.exit("Missing font: %s (PDF needs Arial TTFs; edit FONTS for your OS)." % f)
        pdf.add_font("Arial", style, str(f))
    pdf.set_auto_page_break(True, margin=15)
    pdf.set_margins(18, 16, 18)
    EPW = pdf.epw

    seen_first_h1 = False
    pdf.add_page()
    for typ, payload in tokens:
        if typ == "pagebreak":
            pdf.add_page()
        elif typ.startswith("h"):
            level = int(typ[1])
            if level == 1 and not seen_first_h1:
                seen_first_h1 = True
                pdf.ln(24)
                pdf.set_font("Arial", "B", 30); pdf.set_text_color(31, 51, 100)
                pdf.block(14, clean_heading(payload), align="C")
                pdf.ln(4)
                pdf.set_font("Arial", "", 13); pdf.set_text_color(85, 85, 85)
                pdf.block(8, "Defense Technology Platform  |  Business Plan  |  V4 (Final)", align="C")
                pdf.set_text_color(0, 0, 0)
                continue
            sizes = {1: 17, 2: 13, 3: 11.5, 4: 10.5}
            if level == 1:
                if pdf.get_y() > 40:
                    pdf.add_page()
                pdf.set_text_color(31, 51, 100)
            else:
                pdf.set_text_color(40, 60, 100)
            pdf.set_font("Arial", "B", sizes.get(level, 10.5))
            pdf.ln(2 if level > 1 else 0)
            pdf.block(sizes.get(level, 10.5) * 0.55, clean_heading(payload))
            pdf.set_text_color(0, 0, 0); pdf.ln(1.5)
            if level == 1:
                insert_diagrams_pdf(pdf, diagrams_for_h1(payload), EPW)
        elif typ == "p":
            pdf.set_font("Arial", "", 10)
            pdf.block(5.0, md_for_pdf(payload), markdown=True)
            pdf.ln(1.5)
        elif typ in ("ul", "ol"):
            pdf.set_font("Arial", "", 10)
            for k, (_, item) in enumerate(payload, 1):
                bullet = ("%d. " % k) if typ == "ol" else "- "
                pdf.block(5.0, bullet + md_for_pdf(item), markdown=True)
            pdf.ln(1.5)
        elif typ == "quote":
            pdf.set_font("Arial", "I", 10); pdf.set_text_color(80, 80, 80)
            pdf.block(5.0, md_for_pdf(payload), markdown=True)
            pdf.set_text_color(0, 0, 0); pdf.ln(1.5)
        elif typ == "code":
            pdf.set_font("Arial", "", 9)
            pdf.set_fill_color(244, 244, 244)
            pdf.block(4.5, payload, fill=True)
            pdf.ln(1.5)
        elif typ == "table":
            render_table_pdf(pdf, parse_table(payload))
    pdf.output(str(OUT_PDF))
    return pdf.page_no()


# ============================= VERIFY =============================
def verify_docx():
    """Reopen the generated docx and assert today's canonical facts survived."""
    d = Document(str(OUT_DOCX))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    text = "\n".join(parts)
    imgs = sum(1 for r in d.part.rels.values() if "image" in r.reltype)

    def c(pat):
        return len(re.findall(pat, text, flags=re.IGNORECASE))

    checks = {
        "no 'first platform-native' superlative (D-008)": c(r"first\s+\**platform-native") == 0,
        "Rs.210 Cr MoQ present (D-012)": c(r"Rs\.?\s*210") > 0,
        "Phase 1 framing present (D-009/D-012)": c(r"Phase\s*1") > 0,
        "OQ-C reconciliation note present": c(r"OQ-C") > 0,
        "Rs.21 Cr/system present (D-012)": c(r"21\s*Cr") > 0,
        "all 10 diagrams embedded": imgs == 10,
        "table count >= 50": len(d.tables) >= 50,
    }
    return checks, imgs, len(d.tables)


def main():
    if not MD.exists():
        sys.exit("Missing %s - run build_investor_plan.py first." % MD)
    tokens = tokenize(MD.read_text(encoding="utf-8"))
    build_docx(tokens)
    pages = build_pdf(tokens)
    print("Wrote %s (%.2f MB)" % (OUT_DOCX.name, OUT_DOCX.stat().st_size / 1e6))
    print("Wrote %s (%.2f MB, %d pages)" % (OUT_PDF.name, OUT_PDF.stat().st_size / 1e6, pages))

    checks, imgs, ntables = verify_docx()
    print("Verification (docx):  images=%d  tables=%d" % (imgs, ntables))
    ok = True
    for label, passed in checks.items():
        print("  [%s] %s" % ("PASS" if passed else "FAIL", label))
        ok = ok and passed
    print("ALL CHECKS PASS" if ok else "*** SOME CHECKS FAILED ***")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
